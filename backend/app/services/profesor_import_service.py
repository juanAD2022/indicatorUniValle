from typing import BinaryIO
from decimal import Decimal, InvalidOperation

import openpyxl
from sqlalchemy.orm import Session

from app.models.profesor import Profesor, ProfesorActivo
from app.models.audit_log import AuditLog
from app.schemas.import_schema import (
    ImportError,
    ImportChange,
    ImportPendingRow,
    ImportPreviewResponse,
)

# =============================================================================
# EXPECTED HEADERS - HOJA 1: PROFESORES_ACTIVOS
# =============================================================================

ACTIVOS_HEADERS = [
    "PROFESORES_ACTIVOS",
    "CATEGORIA",
    "CVLAC",
    "ESTADO",
]

ACTIVOS_FIELD_MAP = {
    "PROFESORES_ACTIVOS": "nombre",
    "CATEGORIA": "categoria",
    "CVLAC": "cvlac",
    "ESTADO": "estado",
}

# =============================================================================
# EXPECTED HEADERS - HOJA 2: PERIODO DATA
# =============================================================================

PERIODO_HEADERS = [
    "PERIODO",
    "PLANTA",
    "VARIACION_PLANTA",
    "CONTRATISTAS",
    "VARIACION_CONTRATISTAS",
    "ASISTENTES",
    "VARIACION_ASISTENTES",
    "COMISION",
    "VARIACION_COMISION",
    "HORAS_ADMINISTRATIVO",
    "HORAS_DOCENCIA",
    "HORAS_EXTENSION",
    "HORAS_INVESTIGACION",
]

PERIODO_FIELD_MAP = {
    "PERIODO": "periodo",
    "PLANTA": "planta",
    "VARIACION_PLANTA": "variacion_planta",
    "CONTRATISTAS": "contratistas",
    "VARIACION_CONTRATISTAS": "variacion_contratistas",
    "ASISTENTES": "asistentes",
    "VARIACION_ASISTENTES": "variacion_asistentes",
    "COMISION": "comision",
    "VARIACION_COMISION": "variacion_comision",
    "HORAS_ADMINISTRATIVO": "horas_administrativo",
    "HORAS_DOCENCIA": "horas_docencia",
    "HORAS_EXTENSION": "horas_extension",
    "HORAS_INVESTIGACION": "horas_investigacion",
}

# =============================================================================
# PARSERS
# =============================================================================

def _parse_int(value) -> int:
    if value is None:
        return 0
    try:
        return int(value)
    except (ValueError, TypeError):
        return 0


def _parse_percentage(value) -> Decimal:
    if value is None:
        return Decimal("0")
    s = str(value).strip()
    if s in ("—", "-", "", "N/A", "n/a"):
        return Decimal("0")
    s = s.replace("%", "")
    s = s.replace(",", ".")
    try:
        return Decimal(s)
    except (InvalidOperation, ValueError):
        return Decimal("0")


# =============================================================================
# SHEET PARSING - GENERIC
# =============================================================================

def _parse_generic_sheet(ws, sheet_type: str, expected_headers: list, field_map: dict) -> tuple[list[dict], list[ImportError]]:
    errors: list[ImportError] = []
    rows: list[dict] = []

    headers = [cell.value for cell in ws[1]]
    headers = [h for h in headers if h is not None]
    headers_upper = [str(h).strip().upper() for h in headers]

    missing = [h for h in expected_headers if h not in headers_upper]
    if missing:
        errors.append(
            ImportError(row=1, field="headers", message=f"Faltan columnas en hoja {sheet_type}: {', '.join(missing)}")
        )
        return [], errors

    col_indices = {}
    for i, h in enumerate(headers):
        h_upper = str(h).strip().upper()
        if h_upper in field_map:
            col_indices[h_upper] = i

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(cell is None for cell in row):
            continue

        record = {"_row": row_idx, "_sheet": sheet_type}
        for header, col_idx in col_indices.items():
            if col_idx < len(row):
                record[field_map[header]] = row[col_idx]
        rows.append(record)

    return rows, errors


# =============================================================================
# EXCEL PARSING
# =============================================================================

def parse_profesor_excel(file: BinaryIO) -> tuple[list[dict], list[ImportError]]:
    wb = openpyxl.load_workbook(file, read_only=True)
    all_errors: list[ImportError] = []
    all_rows: list[dict] = []

    for name in wb.sheetnames:
        ws = wb[name]
        headers = [cell.value for cell in ws[1]]
        headers = [str(h).strip().upper() for h in headers if h is not None]

        # Detectar hoja de profesores activos
        if "PROFESORES_ACTIVOS" in name.upper() or all(h in headers for h in ["PROFESORES_ACTIVOS", "CATEGORIA"]):
            rows, errors = _parse_generic_sheet(ws, "PROFESORES_ACTIVOS", ACTIVOS_HEADERS, ACTIVOS_FIELD_MAP)
            all_rows.extend(rows)
            all_errors.extend(errors)
        # Detectar hoja de periodo
        elif "PERIODO" in headers or all(h in headers for h in ["PERIODO", "PLANTA"]):
            rows, errors = _parse_generic_sheet(ws, "PERIODO", PERIODO_HEADERS, PERIODO_FIELD_MAP)
            all_rows.extend(rows)
            all_errors.extend(errors)

    wb.close()

    if not all_rows:
        all_errors.append(
            ImportError(row=0, field="sheets", message="No se encontraron hojas válidas en el archivo")
        )

    return all_rows, all_errors


# =============================================================================
# VALIDATION
# =============================================================================

def _validate_activos_row(row: dict) -> list[ImportError]:
    errors: list[ImportError] = []
    row_num = row["_row"]

    if not row.get("nombre"):
        errors.append(ImportError(row=row_num, field="nombre", message="Nombre del profesor requerido"))

    if not row.get("categoria"):
        errors.append(ImportError(row=row_num, field="categoria", message="Categoría requerida"))

    return errors


def _validate_periodo_row(row: dict) -> list[ImportError]:
    errors: list[ImportError] = []
    row_num = row["_row"]

    if not row.get("periodo"):
        errors.append(ImportError(row=row_num, field="periodo", message="Periodo requerido"))

    int_fields = [
        "planta", "contratistas", "asistentes", "comision",
        "horas_administrativo", "horas_docencia", "horas_extension", "horas_investigacion",
    ]
    for field in int_fields:
        val = row.get(field)
        if val is not None:
            try:
                int(val)
            except (ValueError, TypeError):
                errors.append(ImportError(row=row_num, field=field, message=f"{field} debe ser un entero"))

    pct_fields = [
        "variacion_planta", "variacion_contratistas",
        "variacion_asistentes", "variacion_comision",
    ]
    for field in pct_fields:
        val = row.get(field)
        if val is not None:
            try:
                _parse_percentage(val)
            except (InvalidOperation, ValueError):
                errors.append(ImportError(row=row_num, field=field, message=f"{field} debe ser un porcentaje válido"))

    return errors


# =============================================================================
# NORMALIZATION
# =============================================================================

def _normalize_activos_row(row: dict) -> dict:
    cvlac = str(row.get("cvlac", "")).strip()
    if not cvlac or cvlac.upper() in ("SIN_CVLAC", "SIN CVLAC", "SIN_CvLAC", "—", "-"):
        cvlac = None

    estado = str(row.get("estado", "ACTIVO")).strip().upper()
    if estado not in ("ACTIVO", "INACTIVO"):
        estado = "ACTIVO"

    return {
        "nombre": str(row.get("nombre", "")).strip(),
        "categoria": str(row.get("categoria", "")).strip().upper(),
        "cvlac": cvlac,
        "estado": estado,
    }


def _normalize_periodo_row(row: dict) -> dict:
    return {
        "periodo": str(row.get("periodo", "")),
        "planta": _parse_int(row.get("planta")),
        "variacion_planta": _parse_percentage(row.get("variacion_planta")),
        "contratistas": _parse_int(row.get("contratistas")),
        "variacion_contratistas": _parse_percentage(row.get("variacion_contratistas")),
        "asistentes": _parse_int(row.get("asistentes")),
        "variacion_asistentes": _parse_percentage(row.get("variacion_asistentes")),
        "comision": _parse_int(row.get("comision")),
        "variacion_comision": _parse_percentage(row.get("variacion_comision")),
        "horas_administrativo": _parse_int(row.get("horas_administrativo")),
        "horas_docencia": _parse_int(row.get("horas_docencia")),
        "horas_extension": _parse_int(row.get("horas_extension")),
        "horas_investigacion": _parse_int(row.get("horas_investigacion")),
    }


# =============================================================================
# DETECT CHANGES
# =============================================================================

def _detect_changes(old: dict, new: dict) -> list[str]:
    changed = []
    for key in new:
        if key in ("_row", "_sheet"):
            continue
        old_val = old.get(key)
        new_val = new[key]
        if str(old_val) != str(new_val):
            changed.append(key)
    return changed


# =============================================================================
# PREVIEW
# =============================================================================

def preview_profesor_import(db: Session, file: BinaryIO, filename: str) -> ImportPreviewResponse:
    rows, parse_errors = parse_profesor_excel(file)

    all_errors = list(parse_errors)
    to_create = 0
    to_update = 0
    unchanged = 0
    sample_changes: list[ImportChange] = []
    pending_data: list[ImportPendingRow] = []

    for row in rows:
        sheet_type = row.get("_sheet", "")

        if sheet_type == "PROFESORES_ACTIVOS":
            validation_errors = _validate_activos_row(row)
            if validation_errors:
                all_errors.extend(validation_errors)
                continue

            normalized = _normalize_activos_row(row)
            row_num = row["_row"]

            existing = (
                db.query(ProfesorActivo)
                .filter(
                    ProfesorActivo.nombre == normalized["nombre"],
                    ProfesorActivo.categoria == normalized["categoria"],
                )
                .first()
            )

            if not existing:
                to_create += 1
                pending_data.append(ImportPendingRow(action="create", row=row_num, data=normalized))
                if len(sample_changes) < 10:
                    sample_changes.append(
                        ImportChange(
                            row=row_num,
                            periodo=normalized["nombre"],
                            action="create",
                            fields_changed=list(normalized.keys()),
                            new_values=normalized,
                        )
                    )
            else:
                old_dict = {
                    "nombre": existing.nombre,
                    "categoria": existing.categoria,
                    "cvlac": existing.cvlac,
                    "estado": existing.estado,
                }
                fields_changed = _detect_changes(old_dict, normalized)
                if fields_changed:
                    to_update += 1
                    pending_data.append(
                        ImportPendingRow(
                            action="update",
                            row=row_num,
                            data=normalized,
                            record_id=existing.id,
                            old=old_dict,
                        )
                    )
                    if len(sample_changes) < 10:
                        sample_changes.append(
                            ImportChange(
                                row=row_num,
                                periodo=normalized["nombre"],
                                action="update",
                                fields_changed=fields_changed,
                                old_values={k: old_dict[k] for k in fields_changed},
                                new_values={k: normalized[k] for k in fields_changed},
                            )
                        )
                else:
                    unchanged += 1

        elif sheet_type == "PERIODO":
            validation_errors = _validate_periodo_row(row)
            if validation_errors:
                all_errors.extend(validation_errors)
                continue

            normalized = _normalize_periodo_row(row)
            row_num = row["_row"]

            existing = (
                db.query(Profesor)
                .filter(Profesor.periodo == normalized["periodo"])
                .first()
            )

            if not existing:
                to_create += 1
                pending_data.append(ImportPendingRow(action="create", row=row_num, data=normalized))
                if len(sample_changes) < 10:
                    sample_changes.append(
                        ImportChange(
                            row=row_num,
                            periodo=normalized["periodo"],
                            action="create",
                            fields_changed=list(normalized.keys()),
                            new_values=normalized,
                        )
                    )
            else:
                old_dict = {
                    "periodo": existing.periodo,
                    "planta": existing.planta,
                    "variacion_planta": existing.variacion_planta,
                    "contratistas": existing.contratistas,
                    "variacion_contratistas": existing.variacion_contratistas,
                    "asistentes": existing.asistentes,
                    "variacion_asistentes": existing.variacion_asistentes,
                    "comision": existing.comision,
                    "variacion_comision": existing.variacion_comision,
                    "horas_administrativo": existing.horas_administrativo,
                    "horas_docencia": existing.horas_docencia,
                    "horas_extension": existing.horas_extension,
                    "horas_investigacion": existing.horas_investigacion,
                }
                fields_changed = _detect_changes(old_dict, normalized)
                if fields_changed:
                    to_update += 1
                    pending_data.append(
                        ImportPendingRow(
                            action="update",
                            row=row_num,
                            data=normalized,
                            record_id=existing.id,
                            old=old_dict,
                        )
                    )
                    if len(sample_changes) < 10:
                        sample_changes.append(
                            ImportChange(
                                row=row_num,
                                periodo=normalized["periodo"],
                                action="update",
                                fields_changed=fields_changed,
                                old_values={k: old_dict[k] for k in fields_changed},
                                new_values={k: normalized[k] for k in fields_changed},
                            )
                        )
                else:
                    unchanged += 1

    return ImportPreviewResponse(
        filename=filename,
        total_rows=len(rows),
        to_create=to_create,
        to_update=to_update,
        unchanged=unchanged,
        errors=all_errors,
        sample_changes=sample_changes,
        pending_rows=pending_data,
    )


# =============================================================================
# EXECUTE
# =============================================================================

def execute_profesor_import(db: Session, rows: list[ImportPendingRow], user_id: int, username: str) -> dict:
    created = 0
    updated = 0

    for item in rows:
        action = item.action
        data = item.data.copy()

        if action == "create":
            # Determinar el modelo segun los campos presentes
            if "nombre" in data and "categoria" in data:
                # ProfesorActivo
                new_record = ProfesorActivo(**data)
                db.add(new_record)
                db.flush()
                entity_name = "profesores_activos"
            elif "periodo" in data and "planta" in data:
                # Profesor (datos agregados)
                new_record = Profesor(**data)
                db.add(new_record)
                db.flush()
                entity_name = "profesores"
            else:
                continue

            log = AuditLog(
                user_id=user_id,
                username=username,
                action="IMPORT_CREATE",
                entity=entity_name,
                record_id=new_record.id,
                old_values=None,
                new_values=data,
                import_batch="direct",
            )
            db.add(log)
            created += 1

        elif action == "update":
            record_id = item.record_id
            old_values = item.old

            # Determinar el modelo segun los campos presentes
            if "nombre" in data and "categoria" in data:
                record = db.query(ProfesorActivo).filter(ProfesorActivo.id == record_id).first()
                entity_name = "profesores_activos"
            elif "periodo" in data and "planta" in data:
                record = db.query(Profesor).filter(Profesor.id == record_id).first()
                entity_name = "profesores"
            else:
                continue

            if record:
                for key, value in data.items():
                    setattr(record, key, value)

                log = AuditLog(
                    user_id=user_id,
                    username=username,
                    action="IMPORT_UPDATE",
                    entity=entity_name,
                    record_id=record_id,
                    old_values=old_values,
                    new_values=data,
                    import_batch="direct",
                )
                db.add(log)
                updated += 1

    db.commit()

    return {
        "success": True,
        "created": created,
        "updated": updated,
        "message": f"Importacion completada: {created} creados, {updated} actualizados.",
    }
