from typing import BinaryIO
from decimal import Decimal, InvalidOperation

import openpyxl
from sqlalchemy.orm import Session

from app.models.extension_social import ExtensionSocial
from app.models.audit_log import AuditLog
from app.schemas.import_schema import (
    ImportError,
    ImportChange,
    ImportPendingRow,
    ImportPreviewResponse,
)

# =============================================================================
# EXPECTED HEADERS
# =============================================================================

EXPECTED_HEADERS = [
    "PERIODO",
    "CONFERENCIAS_DICTADAS",
    "CURSOS_OFRECIDOS",
    "DIPLOMADOS_OFRECIDOS",
    "TALLERES_OFRECIDOS",
    "CONSULTORIAS",
    "ASISTENTES",
    "HORAS_OFRECIDAS",
    "PARTICIPACION_ESTUDIANTIL",
    "PARTICIPACION_EGRESADOS",
    "PARTICIPACION_PROFESORES",
    "INGRESO_NETO",
]

# =============================================================================
# FIELD MAP
# =============================================================================

FIELD_MAP = {
    "PERIODO": "periodo",
    "CONFERENCIAS_DICTADAS": "conferencias_dictadas",
    "CURSOS_OFRECIDOS": "cursos_ofrecidos",
    "DIPLOMADOS_OFRECIDOS": "diplomados_ofrecidos",
    "TALLERES_OFRECIDOS": "talleres_ofrecidos",
    "CONSULTORIAS": "consultorias",
    "ASISTENTES": "asistentes",
    "HORAS_OFRECIDAS": "horas_ofrecidas",
    "PARTICIPACION_ESTUDIANTIL": "participacion_estudiantil",
    "PARTICIPACION_EGRESADOS": "participacion_egresados",
    "PARTICIPACION_PROFESORES": "participacion_profesores",
    "INGRESO_NETO": "ingreso_neto",
}

# =============================================================================
# PARSERS
# =============================================================================

def _parse_int(value) -> int:
    if value is None:
        return 0
    try:
        return int(value)
    except (ValueError, TypeError):
        return 0


def _parse_decimal_comma(value) -> Decimal:
    if value is None:
        return Decimal("0")
    s = str(value).strip()
    # Si tiene múltiples puntos, son separadores de miles (ej: 3.850.000)
    if s.count('.') > 1:
        s = s.replace('.', '')
    # Si tiene coma y punto, determinar cuál es separador decimal
    elif ',' in s and '.' in s:
        # Formato 1,234.56 -> quitar comas
        if s.index(',') < s.index('.'):
            s = s.replace(',', '')
        # Formato 1.234,56 -> quitar puntos, coma->punto
        else:
            s = s.replace('.', '').replace(',', '.')
    # Solo coma: determinar si es separador de miles o decimal
    elif ',' in s:
        parts = s.split(',')
        if len(parts) == 2 and len(parts[1]) == 3 and parts[1].isdigit():
            # Formato 420,000 -> separador de miles
            s = s.replace(',', '')
        else:
            # Formato 1234,56 -> separador decimal
            s = s.replace(',', '.')
    try:
        return Decimal(s)
    except (InvalidOperation, ValueError):
        return Decimal("0")


# =============================================================================
# SHEET PARSING
# =============================================================================

def _parse_sheet(ws) -> tuple[list[dict], list[ImportError]]:
    errors: list[ImportError] = []
    rows: list[dict] = []

    headers = [cell.value for cell in ws[1]]
    headers = [h for h in headers if h is not None]
    headers_upper = [str(h).strip().upper() for h in headers]

    missing = [h for h in EXPECTED_HEADERS if h not in headers_upper]
    if missing:
        errors.append(
            ImportError(row=1, field="headers", message=f"Faltan columnas: {', '.join(missing)}")
        )
        return [], errors

    col_indices = {}
    for i, h in enumerate(headers):
        h_upper = str(h).strip().upper()
        if h_upper in FIELD_MAP:
            col_indices[h_upper] = i

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(cell is None for cell in row):
            continue

        record = {"_row": row_idx}
        for header, col_idx in col_indices.items():
            if col_idx < len(row):
                record[FIELD_MAP[header]] = row[col_idx]
        rows.append(record)

    return rows, errors


def parse_extension_social_excel(file: BinaryIO) -> tuple[list[dict], list[ImportError]]:
    wb = openpyxl.load_workbook(file, read_only=True)
    all_errors: list[ImportError] = []
    all_rows: list[dict] = []

    # Buscar la primera hoja que contenga las columnas esperadas
    sheet_found = False
    for name in wb.sheetnames:
        ws = wb[name]
        headers = [cell.value for cell in ws[1]]
        headers = [str(h).strip().upper() for h in headers if h is not None]
        if "PERIODO" in headers:
            rows, errors = _parse_sheet(ws)
            all_rows.extend(rows)
            all_errors.extend(errors)
            sheet_found = True
            break

    if not sheet_found:
        all_errors.append(
            ImportError(row=0, field="sheets", message="No se encontró una hoja con las columnas esperadas")
        )

    wb.close()
    return all_rows, all_errors


# =============================================================================
# VALIDATION
# =============================================================================

def _validate_row(row: dict) -> list[ImportError]:
    errors: list[ImportError] = []
    row_num = row["_row"]

    if not row.get("periodo"):
        errors.append(ImportError(row=row_num, field="periodo", message="Periodo requerido"))

    int_fields = [
        "conferencias_dictadas",
        "cursos_ofrecidos",
        "diplomados_ofrecidos",
        "talleres_ofrecidos",
        "consultorias",
        "asistentes",
        "horas_ofrecidas",
        "participacion_estudiantil",
        "participacion_egresados",
        "participacion_profesores",
    ]
    for field in int_fields:
        val = row.get(field)
        if val is not None:
            try:
                int(val)
            except (ValueError, TypeError):
                errors.append(ImportError(row=row_num, field=field, message=f"{field} debe ser un entero"))

    val = row.get("ingreso_neto")
    if val is not None:
        try:
            _parse_decimal_comma(val)
        except (InvalidOperation, ValueError):
            errors.append(ImportError(row=row_num, field="ingreso_neto", message="INGRESO_NETO invalido"))

    return errors


# =============================================================================
# NORMALIZATION
# =============================================================================

def _normalize_row(row: dict) -> dict:
    return {
        "periodo": str(row.get("periodo", "")),
        "conferencias_dictadas": _parse_int(row.get("conferencias_dictadas")),
        "cursos_ofrecidos": _parse_int(row.get("cursos_ofrecidos")),
        "diplomados_ofrecidos": _parse_int(row.get("diplomados_ofrecidos")),
        "talleres_ofrecidos": _parse_int(row.get("talleres_ofrecidos")),
        "consultorias": _parse_int(row.get("consultorias")),
        "asistentes": _parse_int(row.get("asistentes")),
        "horas_ofrecidas": _parse_int(row.get("horas_ofrecidas")),
        "participacion_estudiantil": _parse_int(row.get("participacion_estudiantil")),
        "participacion_egresados": _parse_int(row.get("participacion_egresados")),
        "participacion_profesores": _parse_int(row.get("participacion_profesores")),
        "ingreso_neto": _parse_decimal_comma(row.get("ingreso_neto")),
    }


# =============================================================================
# DETECT CHANGES
# =============================================================================

def _detect_changes(old: dict, new: dict) -> list[str]:
    changed = []
    for key in new:
        if key in ("_row",):
            continue
        old_val = old.get(key)
        new_val = new[key]
        if str(old_val) != str(new_val):
            changed.append(key)
    return changed


# =============================================================================
# PREVIEW
# =============================================================================

def preview_extension_social_import(db: Session, file: BinaryIO, filename: str) -> ImportPreviewResponse:
    rows, parse_errors = parse_extension_social_excel(file)

    all_errors = list(parse_errors)
    to_create = 0
    to_update = 0
    unchanged = 0
    sample_changes: list[ImportChange] = []
    pending_data: list[ImportPendingRow] = []

    for row in rows:
        validation_errors = _validate_row(row)
        if validation_errors:
            all_errors.extend(validation_errors)
            continue

        normalized = _normalize_row(row)
        row_num = row["_row"]

        existing = (
            db.query(ExtensionSocial)
            .filter(ExtensionSocial.periodo == normalized["periodo"])
            .first()
        )

        if not existing:
            to_create += 1
            pending_data.append(ImportPendingRow(action="create", row=row_num, data=normalized))
            if len(sample_changes) < 10:
                sample_changes.append(
                    ImportChange(
                        row=row_num,
                        periodo=normalized["periodo"],
                        action="create",
                        fields_changed=list(normalized.keys()),
                        new_values=normalized,
                    )
                )
        else:
            old_dict = {
                "periodo": existing.periodo,
                "conferencias_dictadas": existing.conferencias_dictadas,
                "cursos_ofrecidos": existing.cursos_ofrecidos,
                "diplomados_ofrecidos": existing.diplomados_ofrecidos,
                "talleres_ofrecidos": existing.talleres_ofrecidos,
                "consultorias": existing.consultorias,
                "asistentes": existing.asistentes,
                "horas_ofrecidas": existing.horas_ofrecidas,
                "participacion_estudiantil": existing.participacion_estudiantil,
                "participacion_egresados": existing.participacion_egresados,
                "participacion_profesores": existing.participacion_profesores,
                "ingreso_neto": existing.ingreso_neto,
            }
            fields_changed = _detect_changes(old_dict, normalized)
            if fields_changed:
                to_update += 1
                pending_data.append(
                    ImportPendingRow(
                        action="update",
                        row=row_num,
                        data=normalized,
                        record_id=existing.id,
                        old=old_dict,
                    )
                )
                if len(sample_changes) < 10:
                    sample_changes.append(
                        ImportChange(
                            row=row_num,
                            periodo=normalized["periodo"],
                            action="update",
                            fields_changed=fields_changed,
                            old_values={k: old_dict[k] for k in fields_changed},
                            new_values={k: normalized[k] for k in fields_changed},
                        )
                    )
            else:
                unchanged += 1

    return ImportPreviewResponse(
        filename=filename,
        total_rows=len(rows),
        to_create=to_create,
        to_update=to_update,
        unchanged=unchanged,
        errors=all_errors,
        sample_changes=sample_changes,
        pending_rows=pending_data,
    )


# =============================================================================
# EXECUTE
# =============================================================================

def execute_extension_social_import(db: Session, rows: list[ImportPendingRow], user_id: int, username: str) -> dict:
    created = 0
    updated = 0

    for item in rows:
        action = item.action
        data = item.data.copy()

        if action == "create":
            new_record = ExtensionSocial(**data)
            db.add(new_record)
            db.flush()

            log = AuditLog(
                user_id=user_id,
                username=username,
                action="IMPORT_CREATE",
                entity="extension_social",
                record_id=new_record.id,
                old_values=None,
                new_values=data,
                import_batch="direct",
            )
            db.add(log)
            created += 1

        elif action == "update":
            record_id = item.record_id
            old_values = item.old

            record = db.query(ExtensionSocial).filter(ExtensionSocial.id == record_id).first()
            if record:
                for key, value in data.items():
                    setattr(record, key, value)

                log = AuditLog(
                    user_id=user_id,
                    username=username,
                    action="IMPORT_UPDATE",
                    entity="extension_social",
                    record_id=record_id,
                    old_values=old_values,
                    new_values=data,
                    import_batch="direct",
                )
                db.add(log)
                updated += 1

    db.commit()

    return {
        "success": True,
        "created": created,
        "updated": updated,
        "message": f"Importacion completada: {created} creados, {updated} actualizados.",
    }
