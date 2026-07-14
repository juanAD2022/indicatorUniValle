from typing import BinaryIO
from decimal import Decimal, InvalidOperation

import openpyxl
from sqlalchemy.orm import Session

from app.models.gestion_directiva import PlaneacionCurso, ProgramacionCurso, ApoyoEconomico
from app.models.audit_log import AuditLog
from app.schemas.import_schema import (
    ImportError,
    ImportChange,
    ImportPendingRow,
    ImportPreviewResponse,
)

# =============================================================================
# EXPECTED HEADERS
# =============================================================================

PLANEACION_HEADERS = [
    "PERIODO",
    "CURSO",
    "CUPOS_SOLICITADOS",
    "GRUPOS_ESTIMADOS",
    "CUPO_GRUPO",
    "GRUPOS_APROBADOS",
]

PROGRAMACION_HEADERS = [
    "PERIODO",
    "CURSO",
    "GRUPO",
    "PROFESOR",
    "VINCULACION",
    "HORARIO",
    "EDIFICIO",
    "SALON",
    "MATRICULADOS",
    "CANCELARON",
    "APROBARON",
    "REPROBARON",
]

APOYO_HEADERS = [
    "PERIODO",
    "NIVEL",
    "TIPO_EVENTO",
    "LUGAR",
    "PAIS",
    "PARTICIPANTES",
    "ROL",
    "APOYO_ECONOMICO",
]

# =============================================================================
# FIELD MAPS
# =============================================================================

PLANEACION_FIELD_MAP = {
    "PERIODO": "periodo",
    "CURSO": "curso",
    "CUPOS_SOLICITADOS": "cupos_solicitados",
    "GRUPOS_ESTIMADOS": "grupos_estimados",
    "CUPO_GRUPO": "cupo_grupo",
    "GRUPOS_APROBADOS": "grupos_aprobados",
}

PROGRAMACION_FIELD_MAP = {
    "PERIODO": "periodo",
    "CURSO": "curso",
    "GRUPO": "grupo",
    "PROFESOR": "profesor",
    "VINCULACION": "vinculacion",
    "HORARIO": "horario",
    "EDIFICIO": "edificio",
    "SALON": "salon",
    "MATRICULADOS": "matriculados",
    "CANCELARON": "cancelaron",
    "APROBARON": "aprobaron",
    "REPROBARON": "reprobaron",
}

APOYO_FIELD_MAP = {
    "PERIODO": "periodo",
    "NIVEL": "nivel",
    "TIPO_EVENTO": "tipo_evento",
    "LUGAR": "lugar",
    "PAIS": "pais",
    "PARTICIPANTES": "participantes",
    "ROL": "rol",
    "APOYO_ECONOMICO": "apoyo_economico",
}

# =============================================================================
# PARSERS
# =============================================================================

def _parse_int(value) -> int:
    if value is None:
        return 0
    try:
        return int(value)
    except (ValueError, TypeError):
        return 0


def _parse_decimal_comma(value) -> Decimal:
    if value is None:
        return Decimal("0")
    s = str(value).strip()
    # Si tiene múltiples puntos, son separadores de miles (ej: 3.850.000)
    if s.count('.') > 1:
        s = s.replace('.', '')
    # Si tiene coma y punto, determinar cuál es separador decimal
    elif ',' in s and '.' in s:
        # Formato 1,234.56 -> quitar comas
        if s.index(',') < s.index('.'):
            s = s.replace(',', '')
        # Formato 1.234,56 -> quitar puntos, coma->punto
        else:
            s = s.replace('.', '').replace(',', '.')
    # Solo coma: determinar si es separador de miles o decimal
    elif ',' in s:
        parts = s.split(',')
        if len(parts) == 2 and len(parts[1]) == 3 and parts[1].isdigit():
            # Formato 420,000 -> separador de miles
            s = s.replace(',', '')
        else:
            # Formato 1234,56 -> separador decimal
            s = s.replace(',', '.')
    try:
        return Decimal(s)
    except (InvalidOperation, ValueError):
        return Decimal("0")


# =============================================================================
# SHEET PARSING
# =============================================================================

def _parse_sheet(ws, sheet_type: str, expected_headers: list, field_map: dict) -> tuple[list[dict], list[ImportError]]:
    errors: list[ImportError] = []
    rows: list[dict] = []

    headers = [cell.value for cell in ws[1]]
    headers = [h for h in headers if h is not None]
    headers_upper = [str(h).strip().upper() for h in headers]

    missing = [h for h in expected_headers if h not in headers_upper]
    if missing:
        errors.append(
            ImportError(row=1, field="headers", message=f"Faltan columnas en hoja {sheet_type}: {', '.join(missing)}")
        )
        return [], errors

    col_indices = {}
    for i, h in enumerate(headers):
        h_upper = str(h).strip().upper()
        if h_upper in field_map:
            col_indices[h_upper] = i

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(cell is None for cell in row):
            continue

        record = {"_row": row_idx, "_sheet": sheet_type}
        for header, col_idx in col_indices.items():
            if col_idx < len(row):
                record[field_map[header]] = row[col_idx]
        rows.append(record)

    return rows, errors


def parse_gestion_directiva_excel(file: BinaryIO) -> tuple[list[dict], list[ImportError]]:
    wb = openpyxl.load_workbook(file, read_only=True)
    all_errors: list[ImportError] = []
    all_rows: list[dict] = []

    sheet_map = {}
    for name in wb.sheetnames:
        upper = name.upper().strip()
        if "PLANEACION" in upper or "PLANEACION_CURSOS" in upper:
            sheet_map["PLANEACION_CURSOS"] = wb[name]
        elif "PROGRAMACION" in upper or "PROGRAMACION_CURSOS" in upper:
            sheet_map["PROGRAMACION_CURSOS"] = wb[name]
        elif "APOYO" in upper or "APOYO_ECONOMICO" in upper:
            sheet_map["APOYO_ECONOMICO"] = wb[name]

    if "PLANEACION_CURSOS" in sheet_map:
        rows, errors = _parse_sheet(
            sheet_map["PLANEACION_CURSOS"], "PLANEACION_CURSOS", PLANEACION_HEADERS, PLANEACION_FIELD_MAP
        )
        all_rows.extend(rows)
        all_errors.extend(errors)

    if "PROGRAMACION_CURSOS" in sheet_map:
        rows, errors = _parse_sheet(
            sheet_map["PROGRAMACION_CURSOS"], "PROGRAMACION_CURSOS", PROGRAMACION_HEADERS, PROGRAMACION_FIELD_MAP
        )
        all_rows.extend(rows)
        all_errors.extend(errors)

    if "APOYO_ECONOMICO" in sheet_map:
        rows, errors = _parse_sheet(
            sheet_map["APOYO_ECONOMICO"], "APOYO_ECONOMICO", APOYO_HEADERS, APOYO_FIELD_MAP
        )
        all_rows.extend(rows)
        all_errors.extend(errors)

    if not sheet_map:
        all_errors.append(
            ImportError(row=0, field="sheets", message="No se encontraron hojas PLANEACION_CURSOS, PROGRAMACION_CURSOS o APOYO_ECONOMICO en el archivo")
        )

    wb.close()
    return all_rows, all_errors


# =============================================================================
# VALIDATION
# =============================================================================

def _validate_planeacion_row(row: dict) -> list[ImportError]:
    errors: list[ImportError] = []
    row_num = row["_row"]

    if not row.get("periodo"):
        errors.append(ImportError(row=row_num, field="periodo", message="Periodo requerido"))
    if not row.get("curso"):
        errors.append(ImportError(row=row_num, field="curso", message="Curso requerido"))

    for field in ["cupos_solicitados", "grupos_estimados", "cupo_grupo", "grupos_aprobados"]:
        val = row.get(field)
        if val is not None:
            try:
                int(val)
            except (ValueError, TypeError):
                errors.append(ImportError(row=row_num, field=field, message=f"{field} debe ser un entero"))

    return errors


def _validate_programacion_row(row: dict) -> list[ImportError]:
    errors: list[ImportError] = []
    row_num = row["_row"]

    if not row.get("periodo"):
        errors.append(ImportError(row=row_num, field="periodo", message="Periodo requerido"))
    if not row.get("curso"):
        errors.append(ImportError(row=row_num, field="curso", message="Curso requerido"))
    if not row.get("grupo"):
        errors.append(ImportError(row=row_num, field="grupo", message="Grupo requerido"))

    for field in ["matriculados", "cancelaron", "aprobaron", "reprobaron"]:
        val = row.get(field)
        if val is not None:
            try:
                int(val)
            except (ValueError, TypeError):
                errors.append(ImportError(row=row_num, field=field, message=f"{field} debe ser un entero"))

    return errors


def _validate_apoyo_row(row: dict) -> list[ImportError]:
    errors: list[ImportError] = []
    row_num = row["_row"]

    if not row.get("periodo"):
        errors.append(ImportError(row=row_num, field="periodo", message="Periodo requerido"))

    val = row.get("participantes")
    if val is not None:
        try:
            int(val)
        except (ValueError, TypeError):
            errors.append(ImportError(row=row_num, field="participantes", message="PARTICIPANTES debe ser un entero"))

    val = row.get("apoyo_economico")
    if val is not None:
        try:
            _parse_decimal_comma(val)
        except (InvalidOperation, ValueError):
            errors.append(ImportError(row=row_num, field="apoyo_economico", message="APOYO_ECONOMICO invalido"))

    return errors


# =============================================================================
# NORMALIZATION
# =============================================================================

def _normalize_planeacion_row(row: dict) -> dict:
    return {
        "periodo": str(row.get("periodo", "")),
        "curso": str(row.get("curso", "")),
        "cupos_solicitados": _parse_int(row.get("cupos_solicitados")),
        "grupos_estimados": _parse_int(row.get("grupos_estimados")),
        "cupo_grupo": _parse_int(row.get("cupo_grupo")),
        "grupos_aprobados": _parse_int(row.get("grupos_aprobados")),
    }


def _normalize_programacion_row(row: dict) -> dict:
    return {
        "periodo": str(row.get("periodo", "")),
        "curso": str(row.get("curso", "")),
        "grupo": str(row.get("grupo", "")),
        "profesor": str(row.get("profesor", "")) if row.get("profesor") else None,
        "vinculacion": str(row.get("vinculacion", "")) if row.get("vinculacion") else None,
        "horario": str(row.get("horario", "")) if row.get("horario") else None,
        "edificio": str(row.get("edificio", "")) if row.get("edificio") else None,
        "salon": str(row.get("salon", "")) if row.get("salon") else None,
        "matriculados": _parse_int(row.get("matriculados")),
        "cancelaron": _parse_int(row.get("cancelaron")),
        "aprobaron": _parse_int(row.get("aprobaron")),
        "reprobaron": _parse_int(row.get("reprobaron")),
    }


def _normalize_apoyo_row(row: dict) -> dict:
    return {
        "periodo": str(row.get("periodo", "")),
        "nivel": str(row.get("nivel", "")) if row.get("nivel") else None,
        "tipo_evento": str(row.get("tipo_evento", "")) if row.get("tipo_evento") else None,
        "lugar": str(row.get("lugar", "")) if row.get("lugar") else None,
        "pais": str(row.get("pais", "")) if row.get("pais") else None,
        "participantes": _parse_int(row.get("participantes")),
        "rol": str(row.get("rol", "")) if row.get("rol") else None,
        "apoyo_economico": _parse_decimal_comma(row.get("apoyo_economico")),
    }


# =============================================================================
# DETECT CHANGES
# =============================================================================

def _detect_changes(old: dict, new: dict) -> list[str]:
    changed = []
    for key in new:
        if key in ("_row", "_sheet"):
            continue
        old_val = old.get(key)
        new_val = new[key]
        if str(old_val) != str(new_val):
            changed.append(key)
    return changed


# =============================================================================
# PREVIEW
# =============================================================================

def preview_gestion_directiva_import(db: Session, file: BinaryIO, filename: str) -> ImportPreviewResponse:
    rows, parse_errors = parse_gestion_directiva_excel(file)

    all_errors = list(parse_errors)
    to_create = 0
    to_update = 0
    unchanged = 0
    sample_changes: list[ImportChange] = []
    pending_data: list[ImportPendingRow] = []

    for row in rows:
        sheet_type = row.get("_sheet", "")

        if sheet_type == "PLANEACION_CURSOS":
            validation_errors = _validate_planeacion_row(row)
            if validation_errors:
                all_errors.extend(validation_errors)
                continue
            normalized = _normalize_planeacion_row(row)
            row_num = row["_row"]

            existing = (
                db.query(PlaneacionCurso)
                .filter(
                    PlaneacionCurso.periodo == normalized["periodo"],
                    PlaneacionCurso.curso == normalized["curso"],
                )
                .first()
            )

            if not existing:
                to_create += 1
                pending_data.append(ImportPendingRow(action="create", row=row_num, data=normalized))
                if len(sample_changes) < 10:
                    sample_changes.append(
                        ImportChange(
                            row=row_num,
                            periodo=normalized["periodo"],
                            action="create",
                            fields_changed=list(normalized.keys()),
                            new_values=normalized,
                        )
                    )
            else:
                old_dict = {
                    "periodo": existing.periodo,
                    "curso": existing.curso,
                    "cupos_solicitados": existing.cupos_solicitados,
                    "grupos_estimados": existing.grupos_estimados,
                    "cupo_grupo": existing.cupo_grupo,
                    "grupos_aprobados": existing.grupos_aprobados,
                }
                fields_changed = _detect_changes(old_dict, normalized)
                if fields_changed:
                    to_update += 1
                    pending_data.append(
                        ImportPendingRow(
                            action="update",
                            row=row_num,
                            data=normalized,
                            record_id=existing.id,
                            old=old_dict,
                        )
                    )
                    if len(sample_changes) < 10:
                        sample_changes.append(
                            ImportChange(
                                row=row_num,
                                periodo=normalized["periodo"],
                                action="update",
                                fields_changed=fields_changed,
                                old_values={k: old_dict[k] for k in fields_changed},
                                new_values={k: normalized[k] for k in fields_changed},
                            )
                        )
                else:
                    unchanged += 1

        elif sheet_type == "PROGRAMACION_CURSOS":
            validation_errors = _validate_programacion_row(row)
            if validation_errors:
                all_errors.extend(validation_errors)
                continue
            normalized = _normalize_programacion_row(row)
            row_num = row["_row"]

            existing = (
                db.query(ProgramacionCurso)
                .filter(
                    ProgramacionCurso.periodo == normalized["periodo"],
                    ProgramacionCurso.curso == normalized["curso"],
                    ProgramacionCurso.grupo == normalized["grupo"],
                )
                .first()
            )

            if not existing:
                to_create += 1
                pending_data.append(ImportPendingRow(action="create", row=row_num, data=normalized))
                if len(sample_changes) < 10:
                    sample_changes.append(
                        ImportChange(
                            row=row_num,
                            periodo=normalized["periodo"],
                            action="create",
                            fields_changed=list(normalized.keys()),
                            new_values=normalized,
                        )
                    )
            else:
                old_dict = {
                    "periodo": existing.periodo,
                    "curso": existing.curso,
                    "grupo": existing.grupo,
                    "profesor": existing.profesor,
                    "vinculacion": existing.vinculacion,
                    "horario": existing.horario,
                    "edificio": existing.edificio,
                    "salon": existing.salon,
                    "matriculados": existing.matriculados,
                    "cancelaron": existing.cancelaron,
                    "aprobaron": existing.aprobaron,
                    "reprobaron": existing.reprobaron,
                }
                fields_changed = _detect_changes(old_dict, normalized)
                if fields_changed:
                    to_update += 1
                    pending_data.append(
                        ImportPendingRow(
                            action="update",
                            row=row_num,
                            data=normalized,
                            record_id=existing.id,
                            old=old_dict,
                        )
                    )
                    if len(sample_changes) < 10:
                        sample_changes.append(
                            ImportChange(
                                row=row_num,
                                periodo=normalized["periodo"],
                                action="update",
                                fields_changed=fields_changed,
                                old_values={k: old_dict[k] for k in fields_changed},
                                new_values={k: normalized[k] for k in fields_changed},
                            )
                        )
                else:
                    unchanged += 1

        elif sheet_type == "APOYO_ECONOMICO":
            validation_errors = _validate_apoyo_row(row)
            if validation_errors:
                all_errors.extend(validation_errors)
                continue
            normalized = _normalize_apoyo_row(row)
            row_num = row["_row"]

            # Sin unique constraint, siempre se crea
            to_create += 1
            pending_data.append(ImportPendingRow(action="create", row=row_num, data=normalized))
            if len(sample_changes) < 10:
                sample_changes.append(
                    ImportChange(
                        row=row_num,
                        periodo=normalized["periodo"],
                        action="create",
                        fields_changed=list(normalized.keys()),
                        new_values=normalized,
                    )
                )

    return ImportPreviewResponse(
        filename=filename,
        total_rows=len(rows),
        to_create=to_create,
        to_update=to_update,
        unchanged=unchanged,
        errors=all_errors,
        sample_changes=sample_changes,
        pending_rows=pending_data,
    )


# =============================================================================
# EXECUTE
# =============================================================================

def execute_gestion_directiva_import(db: Session, rows: list[ImportPendingRow], user_id: int, username: str) -> dict:
    created = 0
    updated = 0

    for item in rows:
        action = item.action
        data = item.data.copy()

        if action == "create":
            # Determinar el modelo segun los campos presentes
            if "curso" in data and "cupos_solicitados" in data:
                # PlaneacionCurso
                new_record = PlaneacionCurso(**data)
                db.add(new_record)
                db.flush()
                entity_name = "planeacion_cursos"
            elif "curso" in data and "grupo" in data and "profesor" in data:
                # ProgramacionCurso
                new_record = ProgramacionCurso(**data)
                db.add(new_record)
                db.flush()
                entity_name = "programacion_cursos"
            elif "nivel" in data or "tipo_evento" in data or "apoyo_economico" in data:
                # ApoyoEconomico
                new_record = ApoyoEconomico(**data)
                db.add(new_record)
                db.flush()
                entity_name = "apoyos_economicos"
            else:
                continue

            log = AuditLog(
                user_id=user_id,
                username=username,
                action="IMPORT_CREATE",
                entity=entity_name,
                record_id=new_record.id,
                old_values=None,
                new_values=data,
                import_batch="direct",
            )
            db.add(log)
            created += 1

        elif action == "update":
            record_id = item.record_id
            old_values = item.old

            # Determinar el modelo segun los campos presentes
            if "curso" in data and "cupos_solicitados" in data:
                record = db.query(PlaneacionCurso).filter(PlaneacionCurso.id == record_id).first()
                entity_name = "planeacion_cursos"
            elif "curso" in data and "grupo" in data and "profesor" in data:
                record = db.query(ProgramacionCurso).filter(ProgramacionCurso.id == record_id).first()
                entity_name = "programacion_cursos"
            else:
                continue

            if record:
                for key, value in data.items():
                    setattr(record, key, value)

                log = AuditLog(
                    user_id=user_id,
                    username=username,
                    action="IMPORT_UPDATE",
                    entity=entity_name,
                    record_id=record_id,
                    old_values=old_values,
                    new_values=data,
                    import_batch="direct",
                )
                db.add(log)
                updated += 1

    db.commit()

    return {
        "success": True,
        "created": created,
        "updated": updated,
        "message": f"Importacion completada: {created} creados, {updated} actualizados.",
    }
