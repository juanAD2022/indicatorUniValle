from typing import BinaryIO

import openpyxl
from sqlalchemy.orm import Session

from app.models.grupo_inferir_indicator import GrupoInferirIndicator
from app.models.audit_log import AuditLog
from app.schemas.import_schema import (
    ImportError,
    ImportChange,
    ImportPendingRow,
    ImportPreviewResponse,
)

EXPECTED_HEADERS = [
    "PERIODO",
    "PROFESORES_VINCULADOS",
    "JOVENES_INVESTIGADORES",
    "MONITORES",
    "ASISTENTES_INVESTIGACION",
    "PROYECTOS_DESARROLLO",
    "PROYECTOS_FINALIZADOS",
    "PROYECTOS_CANCELADOS",
    "CONVOCATORIAS_INTERNAS",
    "CONVOCATORIAS_EXTERNAS",
    "CONVOCATORIAS_PROFESORALES",
    "FINANCIACION_COLCIENCIAS",
    "FINANCIACION_UNIVALLE",
    "FINANCIACION_OTROS",
    "LINEA_REGRESION",
    "LINEA_BIOESTADISTICA",
    "LINEA_ANALISIS_DATOS",
    "LINEA_CONTROL_ESTADISTICO",
    "INVESTIGADOR_JUNIOR",
    "INVESTIGADOR_ASOCIADO",
    "INVESTIGADOR_SENIOR",
    "PONENCIAS_NACIONALES",
    "PONENCIAS_INTERNACIONALES",
    "REVISTAS_NACIONALES",
    "REVISTAS_INTERNACIONALES",
    "PUBLICACIONES_EVENTOS",
    "LIBROS",
    "PROYECTOS_ID",
    "INFORMES_INVESTIGACION",
    "PARTICIPACION_PREGRADO",
    "PARTICIPACION_ESPECIALIZACION",
    "PARTICIPACION_MAESTRIA",
]

FIELD_MAP = {
    "PERIODO": "periodo",
    "PROFESORES_VINCULADOS": "profesores_vinculados",
    "JOVENES_INVESTIGADORES": "jovenes_investigadores",
    "MONITORES": "monitores",
    "ASISTENTES_INVESTIGACION": "asistentes_investigacion",
    "PROYECTOS_DESARROLLO": "proyectos_desarrollo",
    "PROYECTOS_FINALIZADOS": "proyectos_finalizados",
    "PROYECTOS_CANCELADOS": "proyectos_cancelados",
    "CONVOCATORIAS_INTERNAS": "convocatorias_internas",
    "CONVOCATORIAS_EXTERNAS": "convocatorias_externas",
    "CONVOCATORIAS_PROFESORALES": "convocatorias_profesorales",
    "FINANCIACION_COLCIENCIAS": "financiacion_colciencias",
    "FINANCIACION_UNIVALLE": "financiacion_univalle",
    "FINANCIACION_OTROS": "financiacion_otros",
    "LINEA_REGRESION": "linea_regresion",
    "LINEA_BIOESTADISTICA": "linea_bioestadistica",
    "LINEA_ANALISIS_DATOS": "linea_analisis_datos",
    "LINEA_CONTROL_ESTADISTICO": "linea_control_estadistico",
    "INVESTIGADOR_JUNIOR": "investigador_junior",
    "INVESTIGADOR_ASOCIADO": "investigador_asociado",
    "INVESTIGADOR_SENIOR": "investigador_senior",
    "PONENCIAS_NACIONALES": "ponencias_nacionales",
    "PONENCIAS_INTERNACIONALES": "ponencias_internacionales",
    "REVISTAS_NACIONALES": "revistas_nacionales",
    "REVISTAS_INTERNACIONALES": "revistas_internacionales",
    "PUBLICACIONES_EVENTOS": "publicaciones_eventos",
    "LIBROS": "libros",
    "PROYECTOS_ID": "proyectos_id",
    "INFORMES_INVESTIGACION": "informes_investigacion",
    "PARTICIPACION_PREGRADO": "participacion_pregrado",
    "PARTICIPACION_ESPECIALIZACION": "participacion_especializacion",
    "PARTICIPACION_MAESTRIA": "participacion_maestria",
}

INTEGER_FIELDS = [
    "profesores_vinculados", "jovenes_investigadores", "monitores",
    "asistentes_investigacion", "proyectos_desarrollo", "proyectos_finalizados",
    "proyectos_cancelados", "convocatorias_internas", "convocatorias_externas",
    "convocatorias_profesorales", "financiacion_colciencias", "financiacion_univalle",
    "financiacion_otros", "linea_regresion", "linea_bioestadistica",
    "linea_analisis_datos", "linea_control_estadistico", "investigador_junior",
    "investigador_asociado", "investigador_senior", "ponencias_nacionales",
    "ponencias_internacionales", "revistas_nacionales", "revistas_internacionales",
    "publicaciones_eventos", "libros", "proyectos_id", "informes_investigacion",
    "participacion_pregrado", "participacion_especializacion", "participacion_maestria",
]


def _parse_int(value) -> int:
    if value is None:
        return 0
    try:
        return int(value)
    except (ValueError, TypeError):
        return 0


def parse_grupo_inferir_excel(file: BinaryIO) -> tuple[list[dict], list[ImportError]]:
    errors: list[ImportError] = []
    rows: list[dict] = []

    wb = openpyxl.load_workbook(file, read_only=True)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    headers = [h for h in headers if h is not None]

    missing = [h for h in EXPECTED_HEADERS if h not in headers]
    if missing:
        errors.append(
            ImportError(row=1, field="headers", message=f"Faltan columnas: {', '.join(missing)}")
        )
        wb.close()
        return [], errors

    col_indices = {h: i for i, h in enumerate(headers) if h in FIELD_MAP}

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(cell is None for cell in row):
            continue

        record = {"_row": row_idx}
        for header, col_idx in col_indices.items():
            if col_idx < len(row):
                record[FIELD_MAP[header]] = row[col_idx]
        rows.append(record)

    wb.close()
    return rows, errors


def _validate_grupo_inferir_row(row: dict) -> list[ImportError]:
    errors: list[ImportError] = []
    row_num = row["_row"]

    if not row.get("periodo"):
        errors.append(ImportError(row=row_num, field="periodo", message="Periodo requerido"))

    for field in INTEGER_FIELDS:
        val = row.get(field)
        if val is not None:
            try:
                int(val)
            except (ValueError, TypeError):
                errors.append(ImportError(row=row_num, field=field, message=f"{field} debe ser un entero"))

    return errors


def _normalize_grupo_inferir_row(row: dict) -> dict:
    result = {"periodo": str(row.get("periodo", ""))}
    for field in INTEGER_FIELDS:
        result[field] = _parse_int(row.get(field))
    return result


def _detect_changes(old: dict, new: dict) -> list[str]:
    changed = []
    for key in new:
        if key == "_row":
            continue
        old_val = old.get(key)
        new_val = new[key]
        if str(old_val) != str(new_val):
            changed.append(key)
    return changed


def preview_grupo_inferir_import(db: Session, file: BinaryIO, filename: str) -> ImportPreviewResponse:
    rows, parse_errors = parse_grupo_inferir_excel(file)

    all_errors = list(parse_errors)
    to_create = 0
    to_update = 0
    unchanged = 0
    sample_changes: list[ImportChange] = []
    pending_data: list[ImportPendingRow] = []

    for row in rows:
        validation_errors = _validate_grupo_inferir_row(row)
        if validation_errors:
            all_errors.extend(validation_errors)
            continue

        normalized = _normalize_grupo_inferir_row(row)
        row_num = row["_row"]

        existing = (
            db.query(GrupoInferirIndicator)
            .filter(GrupoInferirIndicator.periodo == normalized["periodo"])
            .first()
        )

        if not existing:
            to_create += 1
            pending_data.append(ImportPendingRow(action="create", row=row_num, data=normalized))
            if len(sample_changes) < 10:
                sample_changes.append(
                    ImportChange(
                        row=row_num,
                        periodo=normalized["periodo"],
                        action="create",
                        fields_changed=list(normalized.keys()),
                        new_values=normalized,
                    )
                )
        else:
            old_dict = {col: getattr(existing, col) for col in INTEGER_FIELDS}
            old_dict["periodo"] = existing.periodo

            fields_changed = _detect_changes(old_dict, normalized)
            if fields_changed:
                to_update += 1
                pending_data.append(
                    ImportPendingRow(
                        action="update",
                        row=row_num,
                        data=normalized,
                        record_id=existing.id,
                        old=old_dict,
                    )
                )
                if len(sample_changes) < 10:
                    sample_changes.append(
                        ImportChange(
                            row=row_num,
                            periodo=normalized["periodo"],
                            action="update",
                            fields_changed=fields_changed,
                            old_values={k: old_dict[k] for k in fields_changed},
                            new_values={k: normalized[k] for k in fields_changed},
                        )
                    )
            else:
                unchanged += 1

    return ImportPreviewResponse(
        filename=filename,
        total_rows=len(rows),
        to_create=to_create,
        to_update=to_update,
        unchanged=unchanged,
        errors=all_errors,
        sample_changes=sample_changes,
        pending_rows=pending_data,
    )


def execute_grupo_inferir_import(db: Session, rows: list[ImportPendingRow], user_id: int, username: str) -> dict:
    created = 0
    updated = 0

    for item in rows:
        action = item.action
        data = item.data.copy()

        if action == "create":
            new_record = GrupoInferirIndicator(**data)
            db.add(new_record)
            db.flush()

            log = AuditLog(
                user_id=user_id,
                username=username,
                action="IMPORT_CREATE",
                entity="grupo_inferir_indicators",
                record_id=new_record.id,
                old_values=None,
                new_values=data,
                import_batch="direct",
            )
            db.add(log)
            created += 1

        elif action == "update":
            record_id = item.record_id
            old_values = item.old
            record = db.query(GrupoInferirIndicator).filter(GrupoInferirIndicator.id == record_id).first()
            if record:
                for key, value in data.items():
                    setattr(record, key, value)

                log = AuditLog(
                    user_id=user_id,
                    username=username,
                    action="IMPORT_UPDATE",
                    entity="grupo_inferir_indicators",
                    record_id=record_id,
                    old_values=old_values,
                    new_values=data,
                    import_batch="direct",
                )
                db.add(log)
                updated += 1

    db.commit()

    return {
        "success": True,
        "created": created,
        "updated": updated,
        "message": f"Importación completada: {created} creados, {updated} actualizados.",
    }
