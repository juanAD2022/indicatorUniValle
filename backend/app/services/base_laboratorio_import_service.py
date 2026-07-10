from typing import BinaryIO

import openpyxl
from sqlalchemy.orm import Session

from app.models.base_laboratorio import BaseLaboratorio
from app.models.audit_log import AuditLog
from app.schemas.import_schema import (
    ImportError,
    ImportChange,
    ImportPendingRow,
    ImportPreviewResponse,
)

EXPECTED_HEADERS = [
    "PERIODO",
    "USUARIOS_ESTUDIANTES",
    "USUARIOS_PROFESORES",
    "USUARIOS_EXTERNOS",
    "SERVICIOS_SOLICITADOS",
    "SERVICIOS_ATENDIDOS",
    "PROMEDIO_HORAS_USO",
    "SATISFACCION",
    "EQUIPOS_DISPONIBLES",
    "EQUIPOS_FUERA_SERVICIO",
    "INCIDENTES",
]

FIELD_MAP = {
    "PERIODO": "periodo",
    "USUARIOS_ESTUDIANTES": "usuarios_estudiantes",
    "USUARIOS_PROFESORES": "usuarios_profesores",
    "USUARIOS_EXTERNOS": "usuarios_externos",
    "SERVICIOS_SOLICITADOS": "servicios_solicitados",
    "SERVICIOS_ATENDIDOS": "servicios_atendidos",
    "PROMEDIO_HORAS_USO": "promedio_horas_uso",
    "SATISFACCION": "satisfaccion",
    "EQUIPOS_DISPONIBLES": "equipos_disponibles",
    "EQUIPOS_FUERA_SERVICIO": "equipos_fuera_servicio",
    "INCIDENTES": "incidentes",
}

INTEGER_FIELDS = [
    "usuarios_estudiantes", "usuarios_profesores", "usuarios_externos",
    "servicios_solicitados", "servicios_atendidos",
    "equipos_disponibles", "equipos_fuera_servicio", "incidentes",
]

FLOAT_FIELDS = ["promedio_horas_uso", "satisfaccion"]


def _parse_int(value) -> int:
    if value is None:
        return 0
    try:
        return int(value)
    except (ValueError, TypeError):
        return 0


def _parse_float(value) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    str_value = str(value).replace(",", ".")
    try:
        return float(str_value)
    except (ValueError, TypeError):
        return 0.0


def parse_base_laboratorio_excel(file: BinaryIO) -> tuple[list[dict], list[ImportError]]:
    errors: list[ImportError] = []
    rows: list[dict] = []

    wb = openpyxl.load_workbook(file, read_only=True)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    headers = [h for h in headers if h is not None]
    
    # Normalizar headers a mayúsculas para comparación case-insensitive
    headers_upper = [str(h).strip().upper() for h in headers]

    missing = [h for h in EXPECTED_HEADERS if h not in headers_upper]
    if missing:
        errors.append(
            ImportError(row=1, field="headers", message=f"Faltan columnas: {', '.join(missing)}")
        )
        wb.close()
        return [], errors

    # Crear mapa de índices usando headers en mayúsculas
    col_indices = {}
    for i, h in enumerate(headers):
        h_upper = str(h).strip().upper()
        if h_upper in FIELD_MAP:
            col_indices[h_upper] = i

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(cell is None for cell in row):
            continue

        record = {"_row": row_idx}
        for header, col_idx in col_indices.items():
            if col_idx < len(row):
                record[FIELD_MAP[header]] = row[col_idx]
        rows.append(record)

    wb.close()
    return rows, errors


def _validate_base_laboratorio_row(row: dict) -> list[ImportError]:
    errors: list[ImportError] = []
    row_num = row["_row"]

    if not row.get("periodo"):
        errors.append(ImportError(row=row_num, field="periodo", message="Periodo requerido"))

    for field in INTEGER_FIELDS:
        val = row.get(field)
        if val is not None:
            try:
                int(val)
            except (ValueError, TypeError):
                errors.append(ImportError(row=row_num, field=field, message=f"{field} debe ser un entero"))

    for field in FLOAT_FIELDS:
        val = row.get(field)
        if val is not None:
            try:
                str(val).replace(",", ".")
                float(str(val).replace(",", "."))
            except (ValueError, TypeError):
                errors.append(ImportError(row=row_num, field=field, message=f"{field} debe ser un número decimal"))

    return errors


def _normalize_base_laboratorio_row(row: dict) -> dict:
    result = {"periodo": str(row.get("periodo", ""))}
    for field in INTEGER_FIELDS:
        result[field] = _parse_int(row.get(field))
    for field in FLOAT_FIELDS:
        result[field] = _parse_float(row.get(field))
    return result


def _detect_changes(old: dict, new: dict) -> list[str]:
    changed = []
    for key in new:
        if key == "_row":
            continue
        old_val = old.get(key)
        new_val = new[key]
        if str(old_val) != str(new_val):
            changed.append(key)
    return changed


def preview_base_laboratorio_import(db: Session, file: BinaryIO, filename: str) -> ImportPreviewResponse:
    rows, parse_errors = parse_base_laboratorio_excel(file)

    all_errors = list(parse_errors)
    to_create = 0
    to_update = 0
    unchanged = 0
    sample_changes: list[ImportChange] = []
    pending_data: list[ImportPendingRow] = []

    for row in rows:
        validation_errors = _validate_base_laboratorio_row(row)
        if validation_errors:
            all_errors.extend(validation_errors)
            continue

        normalized = _normalize_base_laboratorio_row(row)
        row_num = row["_row"]

        existing = (
            db.query(BaseLaboratorio)
            .filter(BaseLaboratorio.periodo == normalized["periodo"])
            .first()
        )

        if not existing:
            to_create += 1
            pending_data.append(ImportPendingRow(action="create", row=row_num, data=normalized))
            if len(sample_changes) < 10:
                sample_changes.append(
                    ImportChange(
                        row=row_num,
                        periodo=normalized["periodo"],
                        action="create",
                        fields_changed=list(normalized.keys()),
                        new_values=normalized,
                    )
                )
        else:
            old_dict = {col: getattr(existing, col) for col in INTEGER_FIELDS + FLOAT_FIELDS}
            old_dict["periodo"] = existing.periodo

            fields_changed = _detect_changes(old_dict, normalized)
            if fields_changed:
                to_update += 1
                pending_data.append(
                    ImportPendingRow(
                        action="update",
                        row=row_num,
                        data=normalized,
                        record_id=existing.id,
                        old=old_dict,
                    )
                )
                if len(sample_changes) < 10:
                    sample_changes.append(
                        ImportChange(
                            row=row_num,
                            periodo=normalized["periodo"],
                            action="update",
                            fields_changed=fields_changed,
                            old_values={k: old_dict[k] for k in fields_changed},
                            new_values={k: normalized[k] for k in fields_changed},
                        )
                    )
            else:
                unchanged += 1

    return ImportPreviewResponse(
        filename=filename,
        total_rows=len(rows),
        to_create=to_create,
        to_update=to_update,
        unchanged=unchanged,
        errors=all_errors,
        sample_changes=sample_changes,
        pending_rows=pending_data,
    )


def execute_base_laboratorio_import(db: Session, rows: list[ImportPendingRow], user_id: int, username: str) -> dict:
    created = 0
    updated = 0

    for item in rows:
        action = item.action
        data = item.data.copy()

        if action == "create":
            new_record = BaseLaboratorio(**data)
            db.add(new_record)
            db.flush()

            log = AuditLog(
                user_id=user_id,
                username=username,
                action="IMPORT_CREATE",
                entity="base_laboratorios",
                record_id=new_record.id,
                old_values=None,
                new_values=data,
                import_batch="direct",
            )
            db.add(log)
            created += 1

        elif action == "update":
            record_id = item.record_id
            old_values = item.old
            record = db.query(BaseLaboratorio).filter(BaseLaboratorio.id == record_id).first()
            if record:
                for key, value in data.items():
                    setattr(record, key, value)

                log = AuditLog(
                    user_id=user_id,
                    username=username,
                    action="IMPORT_UPDATE",
                    entity="base_laboratorios",
                    record_id=record_id,
                    old_values=old_values,
                    new_values=data,
                    import_batch="direct",
                )
                db.add(log)
                updated += 1

    db.commit()

    return {
        "success": True,
        "created": created,
        "updated": updated,
        "message": f"Importación completada: {created} creados, {updated} actualizados.",
    }
