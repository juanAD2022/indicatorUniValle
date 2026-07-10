from typing import BinaryIO
from decimal import Decimal, InvalidOperation

import openpyxl
from sqlalchemy.orm import Session

from app.models.posgrado_indicator import PosgradoIndicator
from app.models.audit_log import AuditLog
from app.schemas.import_schema import (
    ImportError,
    ImportChange,
    ImportPendingRow,
    ImportPreviewResponse,
)

EXPECTED_HEADERS_BASE = [
    "PERIODO",
    "MATRICULADOS",
    "GRADUADOS",
    "DESERTORES",
    "PROMEDIO_ACUMULADO",
    "HOMBRES",
    "MUJERES",
    "EMPLEADOS",
    "DESEMPLEADOS",
    "FINANCIACION_PROPIA",
    "FINANCIACION_BECA",
    "FINANCIACION_CREDITO",
    "PONENCIAS",
    "PUBLICACIONES",
]

EXTRA_HEADERS_MAESTRIA = [
    "PROYECTOS_EN DESARROLLO",
    "PROYECTOS_FINALIZADOS",
]

ALL_HEADERS_MAESTRIA = EXPECTED_HEADERS_BASE + EXTRA_HEADERS_MAESTRIA

FIELD_MAP_BASE = {
    "PERIODO": "periodo",
    "MATRICULADOS": "matriculados",
    "GRADUADOS": "graduados",
    "DESERTORES": "desertores",
    "PROMEDIO_ACUMULADO": "promedio_acumulado",
    "HOMBRES": "hombres",
    "MUJERES": "mujeres",
    "EMPLEADOS": "empleados",
    "DESEMPLEADOS": "desempleados",
    "FINANCIACION_PROPIA": "financiacion_propia",
    "FINANCIACION_BECA": "financiacion_beca",
    "FINANCIACION_CREDITO": "financiacion_credito",
    "PONENCIAS": "ponencias",
    "PUBLICACIONES": "publicaciones",
    "PROYECTOS_EN DESARROLLO": "proyectos_desarrollo",
    "PROYECTOS_FINALIZADOS": "proyectos_finalizados",
}


def _parse_decimal_comma(value) -> Decimal:
    if value is None:
        return Decimal("0")
    s = str(value).strip().replace(",", ".")
    try:
        return Decimal(s)
    except (InvalidOperation, ValueError):
        return Decimal("0")


def _parse_int(value) -> int:
    if value is None:
        return 0
    try:
        return int(value)
    except (ValueError, TypeError):
        return 0


def parse_posgrado_sheet(ws, tipo_programa: str) -> tuple[list[dict], list[ImportError]]:
    errors: list[ImportError] = []
    rows: list[dict] = []

    headers = [cell.value for cell in ws[1]]
    headers = [h for h in headers if h is not None]

    if tipo_programa == "MAESTRIA":
        expected = ALL_HEADERS_MAESTRIA
    else:
        expected = EXPECTED_HEADERS_BASE

    missing = [h for h in expected if h not in headers]
    if missing:
        errors.append(
            ImportError(row=1, field="headers", message=f"Faltan columnas en hoja {tipo_programa}: {', '.join(missing)}")
        )
        return [], errors

    col_indices = {h: i for i, h in enumerate(headers) if h in FIELD_MAP_BASE}

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(cell is None for cell in row):
            continue

        record = {"_row": row_idx, "_tipo_programa": tipo_programa}
        for header, col_idx in col_indices.items():
            if col_idx < len(row):
                record[FIELD_MAP_BASE[header]] = row[col_idx]
        rows.append(record)

    return rows, errors


def parse_posgrado_excel(file: BinaryIO) -> tuple[list[dict], list[ImportError]]:
    wb = openpyxl.load_workbook(file, read_only=True)
    all_errors: list[ImportError] = []
    all_rows: list[dict] = []

    sheet_map = {}
    for name in wb.sheetnames:
        upper = name.upper().strip()
        if "MAESTRIA" in upper or "MAESTR" in upper:
            sheet_map["MAESTRIA"] = wb[name]
        elif "ESPECIALIZACION" in upper or "ESPECIALIZ" in upper:
            sheet_map["ESPECIALIZACION"] = wb[name]

    if not sheet_map:
        all_errors.append(
            ImportError(row=0, field="sheets", message="No se encontraron hojas MAESTRIA o ESPECIALIZACION en el archivo")
        )
        wb.close()
        return [], all_errors

    for tipo, ws in sheet_map.items():
        rows, errors = parse_posgrado_sheet(ws, tipo)
        all_rows.extend(rows)
        all_errors.extend(errors)

    wb.close()
    return all_rows, all_errors


def _validate_posgrado_row(row: dict) -> list[ImportError]:
    errors: list[ImportError] = []
    row_num = row["_row"]
    tipo = row.get("_tipo_programa", "")

    if not row.get("periodo"):
        errors.append(ImportError(row=row_num, field="periodo", message="Periodo requerido"))

    mat = row.get("matriculados")
    if mat is not None:
        try:
            int(mat)
        except (ValueError, TypeError):
            errors.append(ImportError(row=row_num, field="matriculados", message="Matriculados debe ser un entero"))

    grad = row.get("graduados")
    if grad is not None:
        try:
            int(grad)
        except (ValueError, TypeError):
            errors.append(ImportError(row=row_num, field="graduados", message="Graduados debe ser un entero"))

    des = row.get("desertores")
    if des is not None:
        try:
            int(des)
        except (ValueError, TypeError):
            errors.append(ImportError(row=row_num, field="desertores", message="Desertores debe ser un entero"))

    prom = row.get("promedio_acumulado")
    if prom is not None:
        try:
            s = str(prom).strip().replace(",", ".")
            d = Decimal(s)
            if d < 0 or d > 5:
                errors.append(ImportError(row=row_num, field="promedio_acumulado", message="Promedio debe ser 0-5"))
        except (InvalidOperation, ValueError):
            errors.append(ImportError(row=row_num, field="promedio_acumulado", message="Promedio inválido"))

    for field in ["hombres", "mujeres", "empleados", "desempleados",
                   "financiacion_propia", "financiacion_beca", "financiacion_credito",
                   "ponencias", "publicaciones"]:
        val = row.get(field)
        if val is not None:
            try:
                int(val)
            except (ValueError, TypeError):
                errors.append(ImportError(row=row_num, field=field, message=f"{field} debe ser un entero"))

    if tipo == "MAESTRIA":
        for field in ["proyectos_desarrollo", "proyectos_finalizados"]:
            val = row.get(field)
            if val is not None:
                try:
                    int(val)
                except (ValueError, TypeError):
                    errors.append(ImportError(row=row_num, field=field, message=f"{field} debe ser un entero"))

    return errors


def _normalize_posgrado_row(row: dict) -> dict:
    tipo = row.get("_tipo_programa", "")

    result = {
        "periodo": str(row.get("periodo", "")),
        "tipo_programa": tipo,
        "matriculados": _parse_int(row.get("matriculados")),
        "graduados": _parse_int(row.get("graduados")),
        "desertores": _parse_int(row.get("desertores")),
        "promedio_acumulado": _parse_decimal_comma(row.get("promedio_acumulado")),
        "hombres": _parse_int(row.get("hombres")),
        "mujeres": _parse_int(row.get("mujeres")),
        "empleados": _parse_int(row.get("empleados")),
        "desempleados": _parse_int(row.get("desempleados")),
        "financiacion_propia": _parse_int(row.get("financiacion_propia")),
        "financiacion_beca": _parse_int(row.get("financiacion_beca")),
        "financiacion_credito": _parse_int(row.get("financiacion_credito")),
        "ponencias": _parse_int(row.get("ponencias")),
        "publicaciones": _parse_int(row.get("publicaciones")),
        "proyectos_desarrollo": _parse_int(row.get("proyectos_desarrollo")) if tipo == "MAESTRIA" else None,
        "proyectos_finalizados": _parse_int(row.get("proyectos_finalizados")) if tipo == "MAESTRIA" else None,
    }

    return result


def _detect_changes(old: dict, new: dict) -> list[str]:
    changed = []
    for key in new:
        if key == "_row" or key == "_tipo_programa":
            continue
        old_val = old.get(key)
        new_val = new[key]
        if str(old_val) != str(new_val):
            changed.append(key)
    return changed


def preview_posgrado_import(db: Session, file: BinaryIO, filename: str) -> ImportPreviewResponse:
    rows, parse_errors = parse_posgrado_excel(file)

    all_errors = list(parse_errors)
    to_create = 0
    to_update = 0
    unchanged = 0
    sample_changes: list[ImportChange] = []
    pending_data: list[ImportPendingRow] = []

    for row in rows:
        validation_errors = _validate_posgrado_row(row)
        if validation_errors:
            all_errors.extend(validation_errors)
            continue

        normalized = _normalize_posgrado_row(row)
        row_num = row["_row"]

        existing = (
            db.query(PosgradoIndicator)
            .filter(
                PosgradoIndicator.periodo == normalized["periodo"],
                PosgradoIndicator.tipo_programa == normalized["tipo_programa"],
            )
            .first()
        )

        if not existing:
            to_create += 1
            pending_data.append(ImportPendingRow(action="create", row=row_num, data=normalized))
            if len(sample_changes) < 10:
                sample_changes.append(
                    ImportChange(
                        row=row_num,
                        periodo=normalized["periodo"],
                        action="create",
                        fields_changed=list(normalized.keys()),
                        new_values=normalized,
                    )
                )
        else:
            old_dict = {
                "periodo": existing.periodo,
                "tipo_programa": existing.tipo_programa,
                "matriculados": existing.matriculados,
                "graduados": existing.graduados,
                "desertores": existing.desertores,
                "promedio_acumulado": float(existing.promedio_acumulado),
                "hombres": existing.hombres,
                "mujeres": existing.mujeres,
                "empleados": existing.empleados,
                "desempleados": existing.desempleados,
                "financiacion_propia": existing.financiacion_propia,
                "financiacion_beca": existing.financiacion_beca,
                "financiacion_credito": existing.financiacion_credito,
                "ponencias": existing.ponencias,
                "publicaciones": existing.publicaciones,
                "proyectos_desarrollo": existing.proyectos_desarrollo,
                "proyectos_finalizados": existing.proyectos_finalizados,
            }

            fields_changed = _detect_changes(old_dict, normalized)
            if fields_changed:
                to_update += 1
                pending_data.append(
                    ImportPendingRow(
                        action="update",
                        row=row_num,
                        data=normalized,
                        record_id=existing.id,
                        old=old_dict,
                    )
                )
                if len(sample_changes) < 10:
                    sample_changes.append(
                        ImportChange(
                            row=row_num,
                            periodo=normalized["periodo"],
                            action="update",
                            fields_changed=fields_changed,
                            old_values={k: old_dict[k] for k in fields_changed},
                            new_values={k: normalized[k] for k in fields_changed},
                        )
                    )
            else:
                unchanged += 1

    return ImportPreviewResponse(
        filename=filename,
        total_rows=len(rows),
        to_create=to_create,
        to_update=to_update,
        unchanged=unchanged,
        errors=all_errors,
        sample_changes=sample_changes,
        pending_rows=pending_data,
    )


def execute_posgrado_import(db: Session, rows: list[ImportPendingRow], user_id: int, username: str) -> dict:
    created = 0
    updated = 0

    for item in rows:
        action = item.action
        data = item.data.copy()

        if action == "create":
            new_record = PosgradoIndicator(**data)
            db.add(new_record)
            db.flush()

            log = AuditLog(
                user_id=user_id,
                username=username,
                action="IMPORT_CREATE",
                entity="posgrado_indicators",
                record_id=new_record.id,
                old_values=None,
                new_values=data,
                import_batch="direct",
            )
            db.add(log)
            created += 1

        elif action == "update":
            record_id = item.record_id
            old_values = item.old
            record = db.query(PosgradoIndicator).filter(PosgradoIndicator.id == record_id).first()
            if record:
                for key, value in data.items():
                    setattr(record, key, value)

                log = AuditLog(
                    user_id=user_id,
                    username=username,
                    action="IMPORT_UPDATE",
                    entity="posgrado_indicators",
                    record_id=record_id,
                    old_values=old_values,
                    new_values=data,
                    import_batch="direct",
                )
                db.add(log)
                updated += 1

    db.commit()

    return {
        "success": True,
        "created": created,
        "updated": updated,
        "message": f"Importación completada: {created} creados, {updated} actualizados.",
    }
